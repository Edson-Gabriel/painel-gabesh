import hmac
from datetime import datetime
from flask import Blueprint, jsonify, request, render_template
from flask_login import current_user, login_required
from models import db, Acougueiro, Atendente, Pedido
from io import BytesIO
from flask import send_file
from xml.sax.saxutils import escape as xml_escape
from models import ConfiguracaoTV
from zoneinfo import ZoneInfo
from extensions import socketio
# Bibliotecas usadas para geração de planilhas e relatórios em PDF do módulo.
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

# Blueprint do módulo de pedidos de TV; todas as rotas ficam sob /cliente/pedidos-tv.
bp_tv = Blueprint('pedidos_tv', __name__, url_prefix='/cliente/pedidos-tv')

@bp_tv.route('/api/configuracao', methods=['PUT'])
@login_required
def config_tv():
    config = db.session.scalar(db.select(ConfiguracaoTV).where(ConfiguracaoTV.usuario_id == current_user.id))
    if not config:
        config = ConfiguracaoTV(usuario_id=current_user.id, senha_exclusao="0000")
        db.session.add(config)
        
    payload = request.get_json(silent=True) or {}
    senha_atual = str(payload.get("senha_atual", "")).strip()
    nova_senha = str(payload.get("nova_senha", "")).strip()
    
    if not nova_senha:
        return jsonify({"error": "A nova senha não pode ser vazia."}), 400
        
    senha_correta = config.senha_exclusao if config.senha_exclusao else "0000"
    
    if senha_atual != senha_correta:
        return jsonify({"error": "Senha atual incorreta."}), 403
        
    config.senha_exclusao = nova_senha
    db.session.commit()
    return jsonify({"message": "Senha atualizada!"})
# Telas do módulo: painel principal, cadastro, relatórios e visualização em TV.
@bp_tv.route('/')
@login_required
def index(): 
    return render_template('pedidos_tv/index.html')

@bp_tv.route('/cadastro')
@login_required
def cadastro(): 
    return render_template('pedidos_tv/cadastro.html')

@bp_tv.route('/relatorios')
@login_required
def relatorios(): 
    return render_template('pedidos_tv/relatorios.html')

@bp_tv.route('/tv')
@login_required
def tv(): 
    return render_template('pedidos_tv/tv.html')

# Gestão da equipe: cadastro, ativação e desativação de açougueiros e atendentes.
@bp_tv.route('/api/acougueiros', methods=['GET', 'POST'])
@login_required
def api_acougueiros():
    if request.method == 'GET':
        rows = db.session.scalars(
            db.select(Acougueiro)
            .where(Acougueiro.usuario_id == current_user.id)
            .order_by(Acougueiro.ativo.desc(), Acougueiro.nome)
        ).all()
        return jsonify([{"id": r.id, "nome": r.nome, "ativo": r.ativo} for r in rows])
    
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("nome", "")).strip()
    if not name: 
        return jsonify({"error": "Informe o nome."}), 400
    
    novo = Acougueiro(usuario_id=current_user.id, nome=name, ativo=True)
    db.session.add(novo)
    db.session.commit()
    return jsonify({"id": novo.id, "nome": novo.nome, "ativo": 1}), 201

@bp_tv.route('/api/acougueiros/<int:id_item>', methods=['PATCH'])
@login_required
def patch_acougueiro(id_item):
    payload = request.get_json(silent=True) or {}
    obj = db.session.scalar(db.select(Acougueiro).where(Acougueiro.id == id_item, Acougueiro.usuario_id == current_user.id))
    if not obj: 
        return jsonify({"error": "Não encontrado."}), 404
    
    obj.ativo = bool(payload.get("ativo"))
    db.session.commit()
    return jsonify({"id": obj.id, "ativo": int(obj.ativo)})

@bp_tv.route('/api/atendentes', methods=['GET', 'POST'])
@login_required
def api_atendentes():
    if request.method == 'GET':
        rows = db.session.scalars(
            db.select(Atendente)
            .where(Atendente.usuario_id == current_user.id)
            .order_by(Atendente.ativo.desc(), Atendente.nome)
        ).all()
        return jsonify([{"id": r.id, "nome": r.nome, "ativo": r.ativo} for r in rows])
    
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("nome", "")).strip()
    if not name: 
        return jsonify({"error": "Informe o nome."}), 400
    
    novo = Atendente(usuario_id=current_user.id, nome=name, ativo=True)
    db.session.add(novo)
    db.session.commit()
    return jsonify({"id": novo.id, "nome": novo.nome, "ativo": 1}), 201

@bp_tv.route('/api/atendentes/<int:id_item>', methods=['PATCH'])
@login_required
def patch_atendente(id_item):
    payload = request.get_json(silent=True) or {}
    obj = db.session.scalar(db.select(Atendente).where(Atendente.id == id_item, Atendente.usuario_id == current_user.id))
    if not obj: 
        return jsonify({"error": "Não encontrado."}), 404
    
    obj.ativo = bool(payload.get("ativo"))
    db.session.commit()
    return jsonify({"id": obj.id, "ativo": int(obj.ativo)})

# Operações de pedido: criação, edição, remoção, conclusão e atualização em tempo real do painel.
@bp_tv.route('/api/pedidos', methods=['GET', 'POST'])
@login_required
def api_pedidos():
    if request.method == 'GET':
        pedidos = db.session.scalars(
            db.select(Pedido)
            .where(Pedido.usuario_id == current_user.id, Pedido.status == 'Pendente')
            .order_by(Pedido.data_retirada, Pedido.hora_retirada)
        ).all()
        return jsonify([{
            "id": p.id, "numero_pedido": p.numero_pedido, "cliente": p.nome_cliente,
            "data_retirada": p.data_retirada, "hora_retirada": p.hora_retirada,
            "retirada_em": f"{p.data_retirada}T{p.hora_retirada}",
            "acougueiro": p.acougueiro.nome if p.acougueiro else "",
            "atendente": p.atendente.nome if p.atendente else "",
            "tipo_entrega": p.tipo_entrega
        } for p in pedidos])

    payload = request.get_json(silent=True) or {}
    numero = str(payload.get("numero_pedido", "")).strip()

    # Garante que cada pedido tenha um identificador único dentro da base do cliente.
    if numero:
        pedido_existente = db.session.scalar(
            db.select(Pedido).where(
                Pedido.usuario_id == current_user.id,
                Pedido.numero_pedido == numero
            )
        )
        if pedido_existente:
            return jsonify({"error": f"O número de pedido '{numero}' já está em uso."}), 400

    novo_pedido = Pedido(
        usuario_id=current_user.id,
        numero_pedido=numero,
        nome_cliente=str(payload.get("cliente", "")).strip(),
        data_retirada=str(payload.get("data_retirada", "")).strip(),
        hora_retirada=str(payload.get("hora_retirada", "")).strip(),
        acougueiro_id=payload.get("acougueiro_id"),
        atendente_id=payload.get("atendente_id") or None,
        tipo_entrega=str(payload.get("tipo_entrega", "Retirada")).strip()
    )
    db.session.add(novo_pedido)
    db.session.commit()
    
    # Notifica a tela de exibição em tempo real para refletir a nova fila imediatamente.
    socketio.emit('atualizar_pedidos', room=f"room_{current_user.id}")
    
    return jsonify({"id": novo_pedido.id}), 201

@bp_tv.route('/api/pedidos/<int:order_id>', methods=['PUT', 'DELETE'])
@login_required
def update_delete_pedido(order_id):
    pedido = db.session.scalar(db.select(Pedido).where(Pedido.id == order_id, Pedido.usuario_id == current_user.id, Pedido.status == 'Pendente'))
    if not pedido: 
        return jsonify({"error": "Pedido não encontrado."}), 404

    payload = request.get_json(silent=True) or {}
    
    if request.method == 'DELETE':
        # Valida a senha de liberação antes de apagar o pedido para evitar remoção acidental.
        config = db.session.scalar(db.select(ConfiguracaoTV).where(ConfiguracaoTV.usuario_id == current_user.id))
        senha_correta = config.senha_exclusao if config else "0000"
        
        if not hmac.compare_digest(str(payload.get("senha", "")), senha_correta):
            return jsonify({"error": "Senha de liberação incorreta."}), 403
            
        db.session.delete(pedido)
        db.session.commit()
        socketio.emit('atualizar_pedidos', room=f"room_{current_user.id}")
        return jsonify({"id": order_id, "excluido": True})

    # Atualiza o pedido quando o cliente altera os dados existentes.
    numero_editado = str(payload.get("numero_pedido", "")).strip()

    # Evita duplicidade ao editar um pedido já cadastrado, ignorando o próprio registro atual.
    if numero_editado and numero_editado != pedido.numero_pedido:
        pedido_existente = db.session.scalar(
            db.select(Pedido).where(
                Pedido.usuario_id == current_user.id,
                Pedido.numero_pedido == numero_editado,
                Pedido.id != order_id
            )
        )
        if pedido_existente:
            return jsonify({"error": f"O número de pedido '{numero_editado}' já pertence a outro pedido."}), 400

    pedido.numero_pedido = numero_editado
    pedido.nome_cliente = str(payload.get("cliente", "")).strip()
    pedido.data_retirada = str(payload.get("data_retirada", "")).strip()
    pedido.hora_retirada = str(payload.get("hora_retirada", "")).strip()
    pedido.acougueiro_id = payload.get("acougueiro_id")
    pedido.atendente_id = payload.get("atendente_id") or None
    pedido.tipo_entrega = str(payload.get("tipo_entrega", "Retirada")).strip()
    
    db.session.commit()
    socketio.emit('atualizar_pedidos', room=f"room_{current_user.id}")
    return jsonify({"id": pedido.id})

@bp_tv.route('/api/pedidos/<int:order_id>/concluir', methods=['POST'])
@login_required
def concluir_pedido(order_id):
    pedido = db.session.scalar(db.select(Pedido).where(Pedido.id == order_id, Pedido.usuario_id == current_user.id, Pedido.status == 'Pendente'))
    if not pedido: 
        return jsonify({"error": "Pedido não encontrado ou já concluído."}), 404
    
    pedido.status = 'Concluído'
    pedido.data_hora_conclusao = datetime.now(ZoneInfo("America/Cuiaba")).isoformat(timespec="seconds")
    db.session.commit()
    socketio.emit('atualizar_pedidos', room=f"room_{current_user.id}")
    return jsonify({"id": order_id, "concluido_em": pedido.data_hora_conclusao})

@bp_tv.route('/api/pedidos/ultimo-numero', methods=['GET'])
@login_required
def ultimo_numero_pedido():
    # Busca os números já utilizados para começar a sequência do próximo pedido sem duplicar identificadores.
    pedidos = db.session.scalars(
        db.select(Pedido.numero_pedido).where(Pedido.usuario_id == current_user.id)
    ).all()
    
    maior = 0
    for num in pedidos:
        if num and num.isdigit():
            maior = max(maior, int(num))
            
    return jsonify({"ultimo_numero": maior})

# Geração de relatórios e exportação em PDF/XLSX por período.
def report_period():
    start = request.args.get("inicio", "")
    end = request.args.get("fim", "")
    try:
        datetime.strptime(start, "%Y-%m-%d")
        datetime.strptime(end, "%Y-%m-%d")
    except ValueError:
        return None, None, {"error": "Informe as datas no formato AAAA-MM-DD."}
    if start > end:
        return None, None, {"error": "A data inicial deve ser anterior ou igual à final."}
    return start, end, None

def get_report_data(start, end):
    # Consulta somente os pedidos concluídos do cliente atual no intervalo informado.
    pedidos_concluidos = db.session.scalars(
        db.select(Pedido)
        .where(
            Pedido.usuario_id == current_user.id,
            Pedido.status == 'Concluído',
            # Compara apenas a data extraída do ISO 8601 armazenado no campo de conclusão.
            db.func.substr(Pedido.data_hora_conclusao, 1, 10) >= start,
            db.func.substr(Pedido.data_hora_conclusao, 1, 10) <= end
        )
        .order_by(Pedido.data_hora_conclusao.desc())
    ).all()

    orders_list = []
    grouped_performance = {}

    for p in pedidos_concluidos:
        # Calcula o atraso em minutos para comparar desempenho e prazos de retirada.
        try:
            dt_retirada = datetime.strptime(f"{p.data_retirada} {p.hora_retirada}", "%Y-%m-%d %H:%M")
            dt_conclusao = datetime.strptime(p.data_hora_conclusao, "%Y-%m-%dT%H:%M:%S")
            diferenca = dt_conclusao - dt_retirada
            minutos_atraso = round(diferenca.total_seconds() / 60.0, 1)
        except (ValueError, TypeError):
            minutos_atraso = 0

        late = minutos_atraso > 0

        order_data = {
            "id": p.id,
            "numero_pedido": p.numero_pedido,
            "cliente": p.nome_cliente,
            "data_retirada": p.data_retirada,
            "hora_retirada": p.hora_retirada,
            "concluido_em": p.data_hora_conclusao.replace('T', ' '),
            "tipo_entrega": p.tipo_entrega,
            "acougueiro_id": p.acougueiro.id if p.acougueiro else None,
            "acougueiro": p.acougueiro.nome if p.acougueiro else "Sem Registro",
            "atendente": p.atendente.nome if p.atendente else "",
            "minutos_atraso": minutos_atraso,
            "resultado": "Atrasado" if late else "No prazo",
            "pedido_id": p.id # Alias usado no frontend
        }
        orders_list.append(order_data)

        # Agrupa os indicadores por açougueiro para montar a visão de desempenho da equipe.
        acou_id = order_data["acougueiro_id"]
        group = grouped_performance.setdefault(acou_id, {
            "acougueiro_id": acou_id,
            "acougueiro": order_data["acougueiro"],
            "total": 0,
            "atrasados": 0,
            "no_prazo": 0,
            "pedidos": [],
        })
        group["total"] += 1
        group["atrasados"] += int(late)
        group["no_prazo"] += int(not late)
        group["pedidos"].append(order_data)

    return orders_list, list(grouped_performance.values())

@bp_tv.route('/api/relatorios/pedidos')
@login_required
def api_relatorios_pedidos():
    start, end, error = report_period()
    if error:
        return jsonify(error), 400
    orders, _ = get_report_data(start, end)
    return jsonify(orders)

@bp_tv.route('/api/relatorios/desempenho')
@login_required
def api_relatorios_desempenho():
    start, end, error = report_period()
    if error:
        return jsonify(error), 400
    _, performance = get_report_data(start, end)
    return jsonify(performance)

# Estilização das abas do relatório exportado para Excel ou PDF.
def style_excel_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="14532D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 40)
        sheet.column_dimensions[column[0].column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

def make_excel_report(orders, performance):
    def safe_cell(value):
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    workbook = Workbook()
    completed = workbook.active
    completed.title = "Pedidos"
    completed.append([
        "Número", "Cliente", "Data da retirada", "Horário", "Tipo", "Açougueiro",
        "Atendente", "Concluído em", "Atraso (min)", "Resultado",
    ])
    for order in orders:
        minutes = round(order["minutos_atraso"] or 0, 1)
        completed.append([safe_cell(value) for value in [
            order["numero_pedido"] or order["id"], order["cliente"],
            order["data_retirada"], order["hora_retirada"], order["tipo_entrega"],
            order["acougueiro"], order["atendente"] or "—", order["concluido_em"],
            minutes, "Atrasado" if minutes > 0 else "No prazo",
        ]])
    style_excel_sheet(completed)

    summary = workbook.create_sheet("Desempenho")
    summary.append(["Açougueiro", "Total", "No prazo", "Atrasados"])
    for group in performance:
        summary.append([safe_cell(group["acougueiro"]), group["total"], group["no_prazo"], group["atrasados"]])
    style_excel_sheet(summary)

    details = workbook.create_sheet("Detalhes")
    details.append([
        "Açougueiro", "Número", "Cliente", "Data da retirada", "Horário",
        "Concluído em", "Atraso (min)", "Resultado",
    ])
    for group in performance:
        for order in group["pedidos"]:
            details.append([safe_cell(value) for value in [
                group["acougueiro"], order["numero_pedido"] or order["pedido_id"],
                order["cliente"], order["data_retirada"], order["hora_retirada"],
                order["concluido_em"], round(order["minutos_atraso"] or 0, 1),
                order["resultado"],
            ]])
    style_excel_sheet(details)
    return workbook

def make_pdf_report(orders, performance, start, end, output):
    document = SimpleDocTemplate(
        output, pagesize=landscape(A4), rightMargin=28, leftMargin=28,
        topMargin=28, bottomMargin=28,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportCell", parent=styles["BodyText"], fontSize=7, leading=9))
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=TA_CENTER))
    story = [
        Paragraph("Relatório de pedidos", styles["ReportTitle"]),
        Paragraph(f"Período de {start} a {end}", styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Desempenho por açougueiro", styles["Heading2"]),
    ]
    summary_data = [["Açougueiro", "Total", "No prazo", "Atrasados"]]
    summary_data.extend([
        [group["acougueiro"], group["total"], group["no_prazo"], group["atrasados"]]
        for group in performance
    ] or [["Sem pedidos no período", 0, 0, 0]])
    story.append(Table(summary_data, repeatRows=1, hAlign="LEFT", style=TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14532D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#A8A29E")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F4")]),
        ("PAD", (0, 0), (-1, -1), 5),
    ])))
    story.extend([Spacer(1, 14), Paragraph("Pedidos concluídos", styles["Heading2"])])

    def cell(value):
        return Paragraph(xml_escape(str(value if value is not None else "—")), styles["ReportCell"])

    order_data = [[cell(value) for value in [
        "Número", "Cliente", "Retirada", "Açougueiro", "Atendente", "Conclusão", "Resultado",
    ]]]
    for order in orders:
        minutes = round(order["minutos_atraso"] or 0, 1)
        order_data.append([cell(value) for value in [
            order["numero_pedido"] or order["id"], order["cliente"],
            f"{order['data_retirada']} {order['hora_retirada']}", order["acougueiro"],
            order["atendente"] or "—", order["concluido_em"],
            f"Atrasado ({minutes:.1f} min)" if minutes > 0 else "No prazo",
        ]])
    if not orders:
        order_data.append([cell("Nenhum pedido concluído no período"), *[cell("") for _ in range(6)]])
    story.append(Table(order_data, repeatRows=1, colWidths=[52, 118, 95, 90, 80, 110, 115], style=TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14532D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#A8A29E")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F4")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("PAD", (0, 0), (-1, -1), 4),
    ])))
    document.build(story)

@bp_tv.route('/api/relatorios/exportar/<formato>')
@login_required
def export_report(formato):
    if formato not in ("pdf", "xlsx"):
        return jsonify({"error": "Formato inválido. Use pdf ou xlsx."}), 404
    start, end, error = report_period()
    if error:
        return jsonify(error), 400
    
    orders, performance = get_report_data(start, end)
    file_buffer = BytesIO()
    filename = f"relatorio_pedidos_{start}_{end}.{formato}"
    
    if formato == "xlsx":
        make_excel_report(orders, performance).save(file_buffer)
        mimetype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        make_pdf_report(orders, performance, start, end, file_buffer)
        mimetype = "application/pdf"
    
    file_buffer.seek(0)
    return send_file(file_buffer, as_attachment=True, download_name=filename, mimetype=mimetype)